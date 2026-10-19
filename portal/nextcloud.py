"""Small Nextcloud client: WebDAV (files), OCS (users, groups, shares), Team folders, Activity, Notifications.

Every call uses the given account. Nextcloud itself enforces permissions: if a file shows up in user X's
PROPFIND, the ACL lets X open it.
"""
import io
import xml.etree.ElementTree as ET
from urllib.parse import quote, unquote

import requests

NS = {"d": "DAV:", "oc": "http://owncloud.org/ns", "nc": "http://nextcloud.org/ns"}
PROPFIND_BODY = """<?xml version="1.0" encoding="UTF-8"?>
<d:propfind xmlns:d="DAV:" xmlns:oc="http://owncloud.org/ns" xmlns:nc="http://nextcloud.org/ns">
  <d:prop>
    <d:getlastmodified/><d:getcontentlength/><d:getcontenttype/><d:resourcetype/>
    <oc:fileid/><oc:owner-id/><oc:permissions/><oc:share-types/><nc:mount-type/><oc:size/>
  </d:prop>
</d:propfind>"""
OCS_HDR = {"OCS-APIRequest": "true", "Accept": "application/json"}


class NextcloudError(RuntimeError):
    pass


class Nextcloud:
    def __init__(self, url, admin_user=None, admin_pass=None, timeout=30):
        self.url = url.rstrip("/")
        self.admin = (admin_user, admin_pass) if admin_user and admin_pass else None
        self.timeout = timeout
        self.s = requests.Session()

    # ------------------------------------------------------------------ dasar
    def _ocs(self, method, path, auth, params=None, data=None):
        params = {"format": "json", **(params or {})}
        r = self.s.request(method, f"{self.url}{path}", params=params, data=data, headers=OCS_HDR,
                           auth=auth, timeout=self.timeout)
        if r.status_code == 401:
            raise NextcloudError("login rejected (401)")
        try:
            js = r.json()
        except ValueError:
            raise NextcloudError(f"{method} {path}: HTTP {r.status_code}, not JSON")
        meta = js.get("ocs", {}).get("meta", {})
        if r.status_code >= 400 or meta.get("statuscode") not in (100, 200):
            raise NextcloudError(f"{method} {path}: HTTP {r.status_code} · {meta.get('statuscode')} {meta.get('message', '')}")
        return js["ocs"].get("data")

    def _need_admin(self):
        if not self.admin:
            raise NextcloudError("admin account not configured (NC_ADMIN_USER / NC_ADMIN_PASS)")
        return self.admin

    def dav_url(self, user, path=""):
        return f"{self.url}/remote.php/dav/files/{quote(user)}/{quote(path.lstrip('/'))}"

    # ------------------------------------------------------------------ users & groups
    def check_login(self, user, pw):
        """{'id', 'displayname'} if user+password are correct, else None."""
        try:
            d = self._ocs("GET", "/ocs/v2.php/cloud/user", (user, pw))
        except NextcloudError:
            return None
        return {"id": d.get("id", user), "displayname": d.get("displayname") or d.get("display-name") or user}

    def user_groups(self, user, auth=None):
        d = self._ocs("GET", f"/ocs/v2.php/cloud/users/{quote(user)}/groups", auth or self._need_admin())
        return list(d.get("groups", []))

    def user_info(self, user):
        return self._ocs("GET", f"/ocs/v2.php/cloud/users/{quote(user)}", self._need_admin())

    def remove_from_group(self, user, group):
        self._ocs("DELETE", f"/ocs/v2.php/cloud/users/{quote(user)}/groups", self._need_admin(), data={"groupid": group})

    def group_members(self, group):
        d = self._ocs("GET", f"/ocs/v2.php/cloud/groups/{quote(group)}/users", self._need_admin())
        return list(d.get("users", []))

    # ------------------------------------------------------------------ Team folders (app groupfolders)
    def team_folders(self):
        """{mount_point: {group: permissions}}. The response format differs between versions, so parse loosely."""
        last = None
        for path in ("/index.php/apps/groupfolders/folders", "/apps/groupfolders/folders"):
            try:
                d = self._ocs("GET", path, self._need_admin())
                break
            except NextcloudError as e:
                last = e
        else:
            raise last
        items = d.values() if isinstance(d, dict) else d
        out = {}
        for f in items or []:
            groups = f.get("groups") or {}
            if isinstance(groups, list):
                groups = {g.get("group_id") or g.get("id"): g.get("permissions", 0) for g in groups}
            norm = {}
            for g, v in groups.items():
                norm[g] = v.get("permissions", 0) if isinstance(v, dict) else v
            out[f.get("mount_point")] = norm
        return out

    # ------------------------------------------------------------------ WebDAV
    def propfind(self, user, pw, path=""):
        r = self.s.request("PROPFIND", self.dav_url(user, path), data=PROPFIND_BODY.encode(),
                           headers={"Depth": "1", "Content-Type": "application/xml"}, auth=(user, pw), timeout=self.timeout)
        if r.status_code != 207:
            raise NextcloudError(f"PROPFIND {path!r} as {user}: HTTP {r.status_code}")
        root, out = ET.fromstring(r.content), []
        marker = f"/remote.php/dav/files/{user}"
        for resp in root.findall("d:response", NS):
            href = unquote(resp.findtext("d:href", default="", namespaces=NS))
            rel = href.split(marker, 1)[-1].rstrip("/") or "/"
            props = {}
            for ps in resp.findall("d:propstat", NS):
                if "200" not in (ps.findtext("d:status", default="", namespaces=NS) or ""):
                    continue
                p = ps.find("d:prop", NS)
                for child in p:
                    tag = child.tag.split("}", 1)[-1]
                    if tag == "resourcetype":
                        props["is_dir"] = child.find("d:collection", NS) is not None
                    elif tag == "share-types":
                        props["share_types"] = [int(x.text) for x in child if x.text and x.text.isdigit()]
                    else:
                        props[tag] = (child.text or "").strip()
            out.append({"path": rel, "is_dir": props.get("is_dir", False), "fileid": props.get("fileid"),
                        "owner": props.get("owner-id"), "mount": props.get("mount-type", ""),
                        "share_types": props.get("share_types", []), "permissions": props.get("permissions", ""),
                        "modified": props.get("getlastmodified", ""), "type": props.get("getcontenttype", ""),
                        "size": int(props.get("getcontentlength") or props.get("size") or 0)})
        here = "/" + path.strip("/") if path.strip("/") else "/"
        return [e for e in out if e["path"] != here]

    def walk(self, user, pw, path=""):
        """Every file and folder the user can see (recursive; Depth: infinity is disabled in Nextcloud)."""
        out, todo = [], [path]
        while todo:
            p = todo.pop()
            for e in self.propfind(user, pw, p):
                out.append(e)
                if e["is_dir"]:
                    todo.append(e["path"])
        return out

    def download(self, user, pw, path):
        r = self.s.get(self.dav_url(user, path), auth=(user, pw), timeout=self.timeout)
        if r.status_code != 200:
            raise NextcloudError(f"GET {path!r}: HTTP {r.status_code}")
        return r.content

    def move(self, user, pw, src, dst):
        r = self.s.request("MOVE", self.dav_url(user, src), auth=(user, pw), timeout=self.timeout,
                           headers={"Destination": self.dav_url(user, dst), "Overwrite": "F"})
        if r.status_code not in (201, 204):
            raise NextcloudError(f"MOVE {src!r} -> {dst!r}: HTTP {r.status_code}")

    def mkdirs(self, user, pw, path):
        parts = [x for x in path.strip("/").split("/") if x]
        for i in range(1, len(parts) + 1):
            r = self.s.request("MKCOL", self.dav_url(user, "/".join(parts[:i])), auth=(user, pw), timeout=self.timeout)
            if r.status_code not in (201, 405):
                raise NextcloudError(f"MKCOL {'/'.join(parts[:i])!r}: HTTP {r.status_code}")

    # ------------------------------------------------------------------ share
    SHARES = "/ocs/v2.php/apps/files_sharing/api/v1/shares"

    def shares_created(self, user, pw):
        """Shares CREATED by this user (to users, groups, public links)."""
        return self._ocs("GET", self.SHARES, (user, pw)) or []

    def shares_received(self, user, pw):
        """Shares RECEIVED by this user (directly or through a group)."""
        return self._ocs("GET", self.SHARES, (user, pw), params={"shared_with_me": "true"}) or []

    def delete_share(self, user, pw, share_id):
        self._ocs("DELETE", f"{self.SHARES}/{share_id}", (user, pw))

    # ------------------------------------------------------------------ activity & notifications
    def activity(self, user, pw, limit=500):
        """The user's file activity (Activity app). [] if the app is disabled."""
        try:   # 304 (no activity) / app disabled -> []
            d = self._ocs("GET", "/ocs/v2.php/apps/activity/api/v2/activity/all", (user, pw), params={"limit": limit})
        except NextcloudError:
            return []
        return [a for a in (d or []) if a.get("object_type", "files") == "files"]

    def notify(self, to_user, title, text=""):
        """Admin notification -> the target user's Nextcloud bell (Notifications app, needs the admin account)."""
        self._ocs("POST", f"/ocs/v2.php/apps/notifications/api/v2/admin_notifications/{quote(to_user)}",
                  self._need_admin(), data={"shortMessage": title[:255], "longMessage": text[:4000]})


def file_text(name, data):
    """Extract text from a file. txt/md/csv directly; docx, xlsx and pdf when the library is installed."""
    low = name.lower()
    if low.endswith((".txt", ".md", ".csv")):
        return data.decode("utf-8", errors="replace")
    if low.endswith(".docx"):
        try:
            import docx
            return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
        except Exception:
            return ""
    if low.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for r in ws.iter_rows(values_only=True):
                    rows.append(" | ".join("" if v is None else str(v) for v in r))
            return "\n".join(rows)
        except Exception:
            return ""
    if low.endswith(".pdf"):
        try:
            import pypdf
            return "\n".join((pg.extract_text() or "") for pg in pypdf.PdfReader(io.BytesIO(data)).pages)
        except Exception:
            return ""
    return ""
