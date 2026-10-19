"""Build the seed files (docx / xlsx / md) and the answer keys from seed/content/*.yaml.

Maintainer tool: the generated files are committed, so people running the lab do NOT need this.
    python scripts/build_content.py
"""
import csv
import datetime as dt
import re
import shutil

from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import common as c

OUT = c.SEED / "files"
KEY = c.SEED / "answer_key"


def local_path(d):
    top = d["path"].split("/")[0]
    return OUT / d["path"] if top in c.TEAM_ROOTS else OUT / "_home" / d["owner"] / d["path"]


def when(d):
    return dt.datetime.combine(dt.date.fromisoformat(str(d["date"])), dt.time(9, 30))


def write_docx(d, path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    rows = None
    for line in d["body"].rstrip().splitlines():
        s = line.strip()
        if s.startswith("|"):
            cells = [x.strip() for x in s.strip("|").split("|")]
            if rows is None:
                rows = doc.add_table(rows=0, cols=len(cells))
                rows.style = "Light Grid Accent 1"
            r = rows.add_row().cells
            for i, v in enumerate(cells[:len(r)]):
                r[i].text = v
            continue
        rows = None
        if not s:
            continue
        if s.startswith("# "):
            h = doc.add_heading(s[2:], level=1)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("- [x] ") or s.startswith("- [ ] "):
            doc.add_paragraph(("☑ " if s[3] == "x" else "☐ ") + s[6:])
        elif s.startswith("- "):
            doc.add_paragraph(s[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", s):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", s), style="List Number")
        else:
            doc.add_paragraph(s)
    cp = doc.core_properties
    cp.author = cp.last_modified_by = c.PEOPLE[d["owner"]]["name"]
    cp.title = path.stem
    cp.created = cp.modified = when(d)
    doc.save(path)


def write_xlsx(d, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(d["sheet"]["columns"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for r in d["sheet"]["rows"]:
        ws.append(r)
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(x.value or "")) for x in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 48)
    wb.properties.creator = wb.properties.lastModifiedBy = c.PEOPLE[d["owner"]]["name"]
    wb.properties.created = wb.properties.modified = when(d)
    wb.save(path)


def main():
    if OUT.exists():
        shutil.rmtree(OUT)
    KEY.mkdir(parents=True, exist_ok=True)
    docs, rows = c.load_docs(), []
    seen = set()
    for d in docs:
        assert d["owner"] in c.PEOPLE, d["path"]
        assert d["path"] not in seen or d["path"].split("/")[0] not in c.TEAM_ROOTS, f"duplicate {d['path']}"
        seen.add(d["path"])
        p = local_path(d)
        p.parent.mkdir(parents=True, exist_ok=True)
        ext = p.suffix.lower()
        if ext == ".docx":
            write_docx(d, p)
        elif ext == ".xlsx":
            write_xlsx(d, p)
        else:
            p.write_text(d["body"].rstrip() + "\n", encoding="utf-8")
        need = sorted(c.expand(d["need"], d["date"]))
        unknown = set(need) - set(c.PEOPLE)
        assert not unknown, (d["path"], unknown)
        rows.append({"path": d["path"], "owner": d["owner"], "location": "team" if d["path"].split("/")[0] in c.TEAM_ROOTS else "home",
                     "date": d["date"], "need_to_know": ";".join(need), "case": d.get("case") or "",
                     "local_file": str(p.relative_to(c.SEED))})
    with open(KEY / "files.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)

    mrows, ids = [], set()
    for m in c.load_emails():
        assert m["id"] not in ids, m["id"]
        ids.add(m["id"])
        for t in [m["from"]] + [x for x in (m.get("to") or []) + (m.get("cc") or []) if not str(x).startswith("list:")]:
            assert t in c.PEOPLE, (m["id"], t)
        mrows.append({"id": m["id"], "message_id": c.message_id(m), "date": m["date"], "from": m["from"],
                      "to": ";".join(map(str, m.get("to") or [])), "cc": ";".join(map(str, m.get("cc") or [])),
                      "subject": m["subject"], "delivered_to": ";".join(sorted(c.email_delivery(m))),
                      "need_to_know": ";".join(sorted(c.email_need(m))), "case": m.get("case") or ""})
    with open(KEY / "mail.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(mrows[0]))
        w.writeheader()
        w.writerows(mrows)
    print(f"{len(rows)} files -> {OUT}")
    print(f"{len(mrows)} emails, {sum(len(r['delivered_to'].split(';')) for r in mrows)} mailbox deliveries -> {KEY / 'mail.csv'}")
    excess = sum(len(set(r["delivered_to"].split(";")) - set(r["need_to_know"].split(";"))) for r in mrows)
    print(f"mailbox copies that the recipient does not need: {excess}")


if __name__ == "__main__":
    main()
